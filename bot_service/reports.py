"""Report generation for the Law Minister Bot."""

import logging
import tempfile
from datetime import datetime, timezone, timedelta

from bot_service import database as db

logger = logging.getLogger("lm_bot.reports")


async def generate_summary_excel(days: int = 7) -> str | None:
    """Generate Excel summary of all messages for the last N days.

    Returns filepath of generated .xlsx or None on failure.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    try:
        messages = await db.get_messages(days=days)
        if not messages:
            return None

        ist = timezone(timedelta(hours=5, minutes=30))

        wb = Workbook()
        ws = wb.active
        ws.title = "Message Summary"

        # Header styling
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")

        headers = ["#", "Direction", "From", "To", "Type", "Content", "Category", "Timestamp (IST)"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for i, row in enumerate(messages, 2):
            ts_raw = row.get("timestamp", "")
            try:
                ts_dt = datetime.fromisoformat(str(ts_raw).replace("Z", "+00:00"))
                ts_ist = ts_dt.astimezone(ist).strftime("%d/%m/%Y %I:%M %p")
            except Exception:
                ts_ist = str(ts_raw)

            ws.cell(row=i, column=1, value=i - 1)
            ws.cell(row=i, column=2, value=row.get("direction", ""))
            ws.cell(row=i, column=3, value=row.get("sender", ""))
            ws.cell(row=i, column=4, value=row.get("recipient", ""))
            ws.cell(row=i, column=5, value=row.get("message_type", ""))
            ws.cell(row=i, column=6, value=(row.get("content", "") or "")[:200])
            ws.cell(row=i, column=7, value=row.get("category", ""))
            ws.cell(row=i, column=8, value=ts_ist)

        # Auto-width columns
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

        filepath = tempfile.mktemp(suffix=".xlsx", prefix="lm_summary_")
        wb.save(filepath)
        logger.info(f"Excel summary generated: {filepath} ({len(messages)} rows)")
        return filepath

    except Exception as e:
        logger.error(f"Excel generation failed: {e}")
        return None


async def generate_attendance_excel(date_str: str = "") -> str | None:
    """Generate Excel report of attendance records."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    try:
        records = await db.get_attendance_records(date_str)
        if not records:
            return None

        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance Report"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")

        headers = ["#", "Staff Name", "Date", "Time", "Status"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for i, row in enumerate(records, 2):
            ws.cell(row=i, column=1, value=i - 1)
            ws.cell(row=i, column=2, value=row.get("staff_name", ""))
            ws.cell(row=i, column=3, value=row.get("date", ""))
            ws.cell(row=i, column=4, value=row.get("time", ""))
            ws.cell(row=i, column=5, value=row.get("status", ""))

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 30)

        filepath = tempfile.mktemp(suffix=".xlsx", prefix="lm_attendance_")
        wb.save(filepath)
        return filepath

    except Exception as e:
        logger.error(f"Attendance Excel generation failed: {e}")
        return None
